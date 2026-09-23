from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import asyncio

import pytest
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from fastapi import UploadFile

import knowledge
import main
from pilot_platform import extract_document_text, PilotValidationError


def pdf_bytes(text="Course boundary policy", password=None):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if password:
        writer.encrypt(password)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def pptx_bytes():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "School access"
    table = slide.shapes.add_table(1, 2, 0, Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Travel distance"
    table.cell(0, 1).text = "Two miles"
    slide.notes_slide.notes_text_frame.text = "Compare neighboring zones"
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


def test_pdf_and_powerpoint_content():
    text = extract_document_text("policy.pdf", pdf_bytes())
    assert "[Page 1]" in text and "boundary policy" in text
    text = extract_document_text("lecture.pptx", pptx_bytes())
    for expected in ("[Slide 1]", "School access", "Travel distance", "Two miles", "Compare neighboring zones"):
        assert expected in text


@pytest.mark.parametrize("name,content", [
    ("scanned.pdf", pdf_bytes(text="")),
    ("locked.pdf", pdf_bytes(password="private")),
    ("broken.pptx", b"not a presentation"),
    ("empty.txt", b"  "),
    ("old.ppt", b"old format"),
])
def test_unreadable_files_are_rejected(name, content):
    with pytest.raises(PilotValidationError):
        extract_document_text(name, content)


def test_legacy_uploads_reload_sources_and_validate_batch():
    with TemporaryDirectory() as directory:
        root = Path(directory)
        with patch.object(main, "KNOWLEDGE_DIR", root), patch.object(knowledge, "KNOWLEDGE_DIR", root), patch.object(main, "_check_admin_access"), patch.object(main, "_validate_legacy_course"), patch.object(main, "_reload_course", return_value={}):
            asyncio.run(main.upload_syllabus(
                request=None, course_id="demo",
                file=UploadFile(filename="syllabus.pdf", file=BytesIO(pdf_bytes()))
            ))
            assert "boundary policy" in knowledge.load_syllabus("demo")
            asyncio.run(main.upload_transcripts(
                request=None, course_id="demo", files=[
                    UploadFile(filename="lesson.pdf", file=BytesIO(pdf_bytes())),
                    UploadFile(filename="lesson.pptx", file=BytesIO(pptx_bytes())),
                ]
            ))
            sources = knowledge.load_transcripts("demo")
            assert set(sources) == {"lesson.pdf", "lesson.pptx"}
            assert "Compare neighboring zones" in sources["lesson.pptx"]
            with pytest.raises(PilotValidationError):
                asyncio.run(main.upload_transcripts(
                    request=None, course_id="demo", files=[
                        UploadFile(filename="new.txt", file=BytesIO(b"valid material")),
                        UploadFile(filename="broken.pdf", file=BytesIO(b"broken")),
                    ]
                ))
            assert not (root / "demo/transcripts/new.txt").exists()
